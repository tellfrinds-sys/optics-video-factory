import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open("curriculum_export.json", encoding="utf-8"))
doors = {d["door_uid"]: d for d in data["doors"]}
units = data["units"]
chapters = {c["chapter_uid"]: c for c in data["chapters"]}
lessons = data["lessons"]

units_by_door = {}
for u in units:
    units_by_door.setdefault(u["door_uid"], []).append(u)
chapters_by_unit = {}
for c in data["chapters"]:
    chapters_by_unit.setdefault(c["unit_uid"], []).append(c)
lessons_by_chapter = {}
for l in lessons:
    lessons_by_chapter.setdefault(l["chapter_uid"], []).append(l)

STATUS_AR = {None: "لم يُنتَج بعد", "ready": "جاهز/منشور", "qa_failed": "فشل المراجعة", "draft": "مسودة"}

wb = Workbook()
ws = wb.active
ws.title = "المنهج الكامل"
ws.sheet_view.rightToLeft = True

FONT = "Arial"
header_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F4E78")
door_font = Font(name=FONT, bold=True, size=12, color="FFFFFF")
door_fill = PatternFill("solid", fgColor="2E75B6")
unit_font = Font(name=FONT, bold=True, size=10.5, color="1F4E78")
unit_fill = PatternFill("solid", fgColor="D9E6F2")
chap_font = Font(name=FONT, italic=True, size=10, color="1F4E78")
chap_fill = PatternFill("solid", fgColor="EEF3F9")
cell_font = Font(name=FONT, size=10.5)
input_font = Font(name=FONT, size=10.5, color="0000FF")
input_fill = PatternFill("solid", fgColor="FFFDE7")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = ["الباب", "الوحدة", "الفصل", "رقم الدرس", "اسم الدرس", "الترتيب الحالي",
           "حالة الإنتاج", "الترتيب الجديد (لو عايز تغيّره)", "ملاحظات"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30

r = 2
for door_uid in sorted(doors, key=lambda k: doors[k].get("door_code") or ""):
    d = doors[door_uid]
    ws.cell(row=r, column=1, value=f'{d.get("door_code","")} — {d["title_ar"]}').font = door_font
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).fill = door_fill
        ws.cell(row=r, column=c).border = border
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
    r += 1
    for u in sorted(units_by_door.get(door_uid, []), key=lambda x: x["unit_uid"]):
        ws.cell(row=r, column=2, value=u["title_ar"]).font = unit_font
        for c in range(2, len(headers) + 1):
            ws.cell(row=r, column=c).fill = unit_fill
            ws.cell(row=r, column=c).border = border
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(headers))
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right", vertical="center")
        r += 1
        for ch in sorted(chapters_by_unit.get(u["unit_uid"], []), key=lambda x: x["chapter_uid"]):
            ws.cell(row=r, column=3, value=ch["title_ar"]).font = chap_font
            for c in range(3, len(headers) + 1):
                ws.cell(row=r, column=c).fill = chap_fill
                ws.cell(row=r, column=c).border = border
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=len(headers))
            ws.cell(row=r, column=3).alignment = Alignment(horizontal="right", vertical="center")
            r += 1
            for les in sorted(lessons_by_chapter.get(ch["chapter_uid"], []), key=lambda x: x["sort_order"] or 0):
                ws.cell(row=r, column=4, value=les["lesson_uid"])
                ws.cell(row=r, column=5, value=les["title_ar"])
                ws.cell(row=r, column=6, value=les["sort_order"])
                ws.cell(row=r, column=7, value=STATUS_AR.get(les.get("video_status"), les.get("video_status") or "لم يُنتَج بعد"))
                new_order = ws.cell(row=r, column=8)
                new_order.fill = input_fill
                new_order.font = input_font
                note = ws.cell(row=r, column=9)
                note.fill = input_fill
                note.font = input_font
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(row=r, column=c)
                    if c not in (8, 9):
                        cell.font = cell_font
                    cell.border = border
                    if c in (4, 5, 6, 7):
                        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=(c == 5))
                r += 1

widths = {1: 6, 2: 6, 3: 6, 4: 12, 5: 46, 6: 14, 7: 16, 8: 22, 9: 28}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w

# legend sheet
leg = wb.create_sheet("دليل الاستخدام")
leg.sheet_view.rightToLeft = True
leg_lines = [
    "دليل استخدام ملف المنهج",
    "",
    "الأعمدة الملوّنة بالأصفر (الترتيب الجديد / ملاحظات) هي الوحيدة المخصّصة للتعديل.",
    "باقي الأعمدة (الباب / الوحدة / الفصل / رقم الدرس / اسم الدرس / الترتيب الحالي / حالة الإنتاج) للمرجعية فقط — لا تُعدَّل يدويًا هنا، أي تغيير فعلي في الترتيب أو المحتوى يُطبَّق لاحقًا على قاعدة البيانات (Supabase) بعد اعتمادك.",
    "",
    "حالة الإنتاج:",
    "  لم يُنتَج بعد — الدرس لسه ما اتعملش له فيديو",
    "  جاهز/منشور — الفيديو موجود ومنشور",
    "  فشل المراجعة — اتعمل لكن ما عداش بوابة الجودة",
    "",
    f"الإجمالي: {len(doors)} باب، {len(units)} وحدة، {len(chapters)} فصل، {len(lessons)} درس.",
    "تاريخ التصدير: 2026-09-11 من قاعدة بيانات Supabase الحيّة.",
]
for i, line in enumerate(leg_lines, start=1):
    cell = leg.cell(row=i, column=1, value=line)
    cell.font = Font(name=FONT, bold=(i == 1), size=13 if i == 1 else 11)
    cell.alignment = Alignment(horizontal="right", wrap_text=True)
leg.column_dimensions["A"].width = 100

wb.save("منهج_بوابة_البصريات.xlsx")
print("saved")
